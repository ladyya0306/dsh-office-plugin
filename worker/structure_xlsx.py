"""Range-aware spreadsheet records; no field/value interpretation is asserted."""
import re
from worker.xlsx_value_evidence import evidence_units
from worker.xlsx_prepare import _scalar, PrepareError
from worker.office_worker import (_load_xlsx, _is_merged_cell, _is_formula,
                                  WorkerError, _merged_range)

UNITS = {"元", "千元", "万元", "亿元", "%"}


def read_xlsx(data, role, scope=None,target_mode='fill_blanks'):
    from worker.source_xlsx import read_source
    physical = {n['ref']: n for n in read_source(data, scope)['nodes']}
    wb = _load_xlsx(data, role)
    nodes, slots, blocked = [], [], []
    try:
        selected = None
        if scope:
            from openpyxl.utils.cell import range_boundaries
            if set(scope) != {"sheet", "range"} or scope["sheet"] not in wb.sheetnames:
                raise WorkerError("invalid_scope", "范围必须是存在的sheet和A1:B10")
            try:
                selected = range_boundaries(scope["range"])
            except (TypeError, ValueError) as exc:
                raise WorkerError("invalid_scope", "无效工作表范围") from exc
            if any(x is None or x < 1 for x in selected) or selected[2] > 256 or selected[3] > 10000:
                raise WorkerError("invalid_scope", "范围超过读取限制")
        for ws in wb.worksheets:
            if scope and ws.title != scope["sheet"]:
                continue
            if "!" in ws.title:
                blocked.append({"sheet": ws.title, "reason": "unaddressable_sheet_name"})
                continue
            limits = selected or (1, 1, ws.max_column, ws.max_row)
            minc, minr, maxc, maxr = limits
            for row in ws.iter_rows(min_row=minr, max_row=maxr, min_col=minc, max_col=maxc):
                for cell in row:
                    if _is_merged_cell(cell):
                        continue
                    left = [str(ws.cell(cell.row,c).value)[:200] for c in range(1,cell.column)
                            if ws.cell(cell.row,c).value is not None and not _is_merged_cell(ws.cell(cell.row,c))]
                    above = []
                    for r in range(1,cell.row):
                        upper = ws.cell(r,cell.column)
                        if upper.value is not None:
                            above.append(str(upper.value)[:200])
                        elif _is_merged_cell(upper):
                            merged = next(m for m in ws.merged_cells.ranges if upper.coordinate in m)
                            value = ws.cell(merged.min_row,merged.min_col).value
                            if value is not None:
                                above.append(str(value)[:200])
                    right = ws.cell(cell.row,cell.column+1).value
                    evidence = ([str(right)] if right in UNITS else []) + left[-3:] + above[-3:]
                    units = evidence_units(evidence)
                    if right in UNITS:
                        units.add(right)
                    if '%' in re.sub(r'"[^"]*"|\\.', '', cell.number_format):
                        units.add('%')
                    unit = next(iter(units)) if len(units)==1 else ""
                    ref = f"{ws.title}!{cell.coordinate}"
                    scalar, value_error = cell.value, None
                    value_type, temporal_evidence = None, None
                    raw_record = physical.get(ref, {})
                    try:
                        if scalar is not None:
                            from datetime import date, time, timedelta
                            if isinstance(scalar, (date, time, timedelta)):
                                if raw_record.get('value_error') or not raw_record.get('value_type'):
                                    raise ValueError('原始日期数值无效或未读取，不能用自动转换替代')
                                scalar, value_type = raw_record['value'], raw_record['value_type']
                                temporal_evidence = raw_record['temporal_evidence']
                                unit = ''; units.clear()
                            else:
                                scalar = _scalar(scalar, ref)
                    except (PrepareError, ValueError) as exc:
                        scalar, value_error = str(cell.value), getattr(exc, 'code', 'invalid_temporal_value')
                    record = {"ref": ref, "kind": "cell", "value": scalar,
                              "text": "" if scalar is None else str(scalar),
                              "row_headers": left[-8:], "column_headers": above[-8:],
                              "parent_context": [ws.title] + [str(ws.cell(r,1).value)[:200] for r in range(max(1,cell.row-8),cell.row)
                                                               if ws.cell(r,1).value is not None],
                              "unit": unit, "unit_evidence": evidence,
                              "unit_ambiguous": len(units)>1, "formula": _is_formula(cell),
                              "value_error": value_error, "merged_span": _merged_range(ws,cell.coordinate),
                              'value_type': value_type, 'temporal_evidence': temporal_evidence,
                              'number_format':cell.number_format,
                              "physical_path": ref, "relation_status": "unconfirmed"}
                    record["context_truncated"] = len(left)>8 or len(above)>8
                    blank = scalar is None or isinstance(scalar,str) and scalar.strip() and all(c in "_＿—-□☐ " for c in scalar)
                    writable = (blank or target_mode=='review_values' and not value_error) and not record["formula"] and not (ws.protection.sheet and cell.protection.locked)
                    # Empty cells inside the actual used range are addressable;
                    # neighbouring labels need not fit a preset layout.
                    if scalar is not None or cell.has_style or writable and (left or above):
                        nodes.append(record)
                    if role == "target" and writable and (left or above or cell.has_style):
                        slots.append({**record, "target_ref": ref, "write_kind": "cell",'requires_expected_text':not blank})
                    if len(nodes)>1000 or len(slots)>1000:
                        raise WorkerError("structure_limit", "请缩小范围；记录超过1000，未悄悄截断")
        return {"nodes": nodes, "slots": slots, "capabilities": {
            "read": True, "write": role=="target", "formula_values": False,'target_mode':target_mode,
            "unsupported": blocked, "scope": scope, "complete": scope is None,
            'table_ranges':{sheet.title+'!'+table.name:table.ref for sheet in wb for table in sheet.tables.values()},
            "complete_requested_range":True}}
    finally:
        wb.close()

