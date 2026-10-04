"""Bind explicit long-term labels and typed keys to current XLSX positions."""
from worker.office_worker import WorkerError
from worker.evidence_read import read_evidence
from worker.structured_worker import _read
from worker.source_xlsx import position


def fail(code, message, refs=()):
    raise WorkerError('relationship_' + code, message + ('：' + '、'.join(refs) if refs else ''),
                      {'references': list(refs)})


def known(value, keys, where):
    if not isinstance(value, dict) or set(value) - set(keys):
        fail('rule_invalid', where + '字段不符合长期规则结构')


def label(value):
    if not isinstance(value, str) or not value.strip() or len(value) > 80:
        fail('rule_invalid', '字段名称需1至80字符的明确原文')
    return value


def validate_rule(rule):
    known(rule, ['schema', 'source_sheet', 'target_sheet', 'keys', 'fields', 'table_name'], '规则')
    if rule.get('schema') != 'office_relationship_rule_v1':
        fail('rule_invalid', '长期规则版本需office_relationship_rule_v1')
    for name in ['source_sheet', 'target_sheet']:
        label(rule.get(name))
        if '!' in rule[name]:
            fail('rule_invalid', '工作表名含!，需要改为明确可引用名称')
    keys, fields = rule.get('keys'), rule.get('fields')
    if not isinstance(keys, list) or not 1 <= len(keys) <= 8 or not isinstance(fields, list) or not 1 <= len(fields) <= 32:
        fail('rule_invalid', '需1至8个匹配键和1至32个字段')
    for key in keys:
        known(key, ['source_label', 'target_label'], '匹配键')
        label(key.get('source_label')); label(key.get('target_label'))
    for side in ['source_label', 'target_label']:
        if len({key[side] for key in keys}) != len(keys):
            fail('rule_invalid', '匹配键字段重复')
    targets = set()
    for field in fields:
        known(field, ['target_label', 'source_labels', 'operation', 'source_unit_label', 'target_unit', 'decimals', 'blank_markers', 'absolute_tolerance'], '字段')
        target = label(field.get('target_label'))
        if target in targets:
            fail('rule_invalid', '目标字段重复：' + target)
        targets.add(target)
        sources = field.get('source_labels')
        if not isinstance(sources, list) or not 1 <= len(sources) <= 8:
            fail('rule_invalid', '每项字段需1至8个来源名称')
        for source in sources: label(source)
        if len(set(sources)) != len(sources): fail('rule_invalid', '同一来源字段不能重复参与计算')
        op = field.get('operation', 'copy')
        if op not in {'copy', 'difference', 'ratio', 'blank'} or len(sources) != (2 if op in {'difference', 'ratio'} else 1):
            fail('rule_invalid', '当前长期规则支持复制、两值差、两值比和明确空标记')
        if bool(field.get('source_unit_label')) != bool(field.get('target_unit')):
            fail('rule_invalid', '来源单位字段和目标单位成对填写')
        if field.get('source_unit_label'):
            label(field['source_unit_label'])
            if field['target_unit'] not in {'元', '千元', '万元', '亿元'} or op == 'blank':
                fail('rule_invalid', '金额目标单位需元、千元、万元或亿元')
        decimals = field.get('decimals')
        if 'absolute_tolerance' in field:
            from worker.output_check import ratio_tolerance
            try: ratio_tolerance(field, 'xlsx')
            except ValueError as exc: fail('rule_invalid', str(exc))
        if decimals is not None and (isinstance(decimals, bool) or not isinstance(decimals, int) or not 0 <= decimals <= 10):
            fail('rule_invalid', '小数位需0至10整数')
        if op == 'blank':
            markers = field.get('blank_markers')
            if not isinstance(markers, list) or not 1 <= len(markers) <= 16 or any(not isinstance(v, str) or len(v) > 128 for v in markers):
                fail('rule_invalid', '空值需1至16个明确文字标记，0不作为空值')
    if not {key['target_label'] for key in keys} <= targets:
        fail('rule_invalid', '目标匹配键也需列入输出字段，逐项核对主体和期间')
    if 'table_name' in rule: label(rule['table_name'])
    return rule


def column(ref):
    return ''.join(c for c in ref.rsplit('!', 1)[1] if c.isalpha())


def header(nodes, sheet, labels):
    rows = {}
    for node in nodes:
        if node['ref'].rsplit('!', 1)[0] != sheet: continue
        row, _ = position(node['ref'].rsplit('!', 1)[1])
        if node.get('value') in labels and not node.get('formula') and not node.get('value_error'):
            rows.setdefault(row, {}).setdefault(node['value'], []).append(node)
    candidates = [(row, values) for row, values in rows.items() if set(values) == labels]
    if len(candidates) != 1:
        refs = [node['ref'] for values in rows.values() for group in values.values() for node in group]
        fail('header_ambiguous' if candidates else 'header_missing',
             sheet + '需唯一完整表头，实际找到' + str(len(candidates)) + '行；核对字段名称和表头范围', refs[:32])
    row, values = candidates[0]
    for name, group in values.items():
        if len(group) != 1: fail('header_ambiguous', sheet + '表头字段重复：' + name, [n['ref'] for n in group])
    return row, {name: group[0] for name, group in values.items()}


def record_keys(nodes, sheet, header_row, columns):
    cells = {node['ref']: node for node in nodes}
    rows = sorted({position(ref.rsplit('!', 1)[1])[0] for ref in cells if ref.rsplit('!', 1)[0] == sheet})
    result = {}
    for row in rows:
        if row <= header_row: continue
        refs = [f'{sheet}!{col}{row}' for col in columns]
        values = [cells.get(ref, {}) for ref in refs]
        if all(n.get('value') is None for n in values): continue
        if any(n.get('value') is None or n.get('value') == '' or n.get('formula') or n.get('value_error') for n in values):
            fail('key_invalid', '记录匹配键缺少值、含公式或读取错误', refs)
        key = tuple((n.get('value_type') or type(n['value']).__name__, n['value']) for n in values)
        if key in result:
            fail('key_duplicate', '主体和期间匹配键重复，需要明确区分这两条记录', result[key]['refs'] + refs)
        result[key] = {'row': row, 'refs': refs}
    if not result: fail('records_empty', sheet + '表头之后需要业务记录')
    return result, cells


def bind_relationship(request):
    rule = validate_rule(request.get('rule'))
    items = request.get('sources')
    if not isinstance(items, list) or len(items) != 1 or items[0].get('format') != 'xlsx' or request.get('target', {}).get('format') != 'xlsx':
        fail('input_format', '当前绑定入口使用一份XLSX来源和一份XLSX模板')
    source = read_evidence(items[0], 0, {'requirements': []})
    target = _read(request['target'], 'target', target_mode='review_values')
    if not source['capabilities'].get('read') or not target['capabilities'].get('read'):
        fail('input_unreadable', '需要先取得来源和目标的可读结构')
    source_labels = {key['source_label'] for key in rule['keys']}
    for field in rule['fields']:
        source_labels.update(field['source_labels'])
        if field.get('source_unit_label'): source_labels.add(field['source_unit_label'])
    target_labels = {field['target_label'] for field in rule['fields']}
    sr, sh = header(source['nodes'], rule['source_sheet'], source_labels)
    tr, th = header(target['nodes'], rule['target_sheet'], target_labels)
    source_keys, source_cells = record_keys(source['nodes'], rule['source_sheet'], sr,
                                           [column(sh[k['source_label']]['ref']) for k in rule['keys']])
    target_keys, _ = record_keys(target['nodes'], rule['target_sheet'], tr,
                                 [column(th[k['target_label']]['ref']) for k in rule['keys']])
    requirements, locations = [], []
    writable = {slot['target_ref'] for slot in target['slots']}
    for key, record in target_keys.items():
        matched = source_keys.get(key)
        if not matched:
            fail('record_missing', '来源中找不到模板要求的主体和期间；补齐这条资料', record['refs'])
        source_row, target_row = matched['row'], record['row']
        locations.append({'source_row': source_row, 'target_row': target_row,
                          'source_key_refs': matched['refs'], 'target_key_refs': record['refs']})
        for field in rule['fields']:
            out_ref = f'{rule["target_sheet"]}!{column(th[field["target_label"]]["ref"])}{target_row}'
            if out_ref not in writable:
                fail('target_not_writable', '目标位置受保护、含公式或类型错误；使用可填模板处理这项', [out_ref])
            op = field.get('operation', 'copy')
            item = {'target_ref': out_ref, 'operation': op}
            value_refs = [f's0::{rule["source_sheet"]}!{column(sh[name]["ref"])}{source_row}' for name in field['source_labels']]
            for ref in value_refs:
                node = source_cells.get(ref.split('::', 1)[1])
                if node is None or node.get('formula') or node.get('value_error'):
                    fail('value_invalid', '来源字段缺少可用值、含公式或读取错误；补齐这项来源', [ref])
            if op == 'blank':
                item.update(source_refs=value_refs, blank_markers=field['blank_markers'])
            else:
                item['source_matches'] = [{'sheet': rule['source_sheet'], 'value_column': column(sh[name]['ref']), 'start_row': sr+1,
                    'key_columns': [{'column': column(sh[k['source_label']]['ref']),
                                     'target_ref': f'{rule["target_sheet"]}!{column(th[k["target_label"]]["ref"])}{target_row}'} for k in rule['keys']]} for name in field['source_labels']]
            for name in ['decimals', 'absolute_tolerance']:
                if name in field: item[name] = field[name]
            if field.get('source_unit_label'):
                unit_ref = f'{rule["source_sheet"]}!{column(sh[field["source_unit_label"]]["ref"])}{source_row}'
                unit_node = source_cells.get(unit_ref, {})
                unit = unit_node.get('value')
                if unit_node.get('formula') or unit_node.get('value_error') or unit not in {'元', '千元', '万元', '亿元'}:
                    fail('unit_invalid', '本条记录单位需明确金额单位', [unit_ref])
                item.update(source_unit=unit, target_unit=field['target_unit'], source_unit_ref='s0::'+unit_ref)
            requirements.append(item)
    if len(requirements) > 10000: fail('coverage_limit', '当前绑定超过10000个字段，需要缩小任务范围')
    contract = {'coverage': [r['target_ref'] for r in requirements], 'requirements': requirements, 'preserve_template': True}
    if rule.get('table_name'):
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        identity = rule['target_sheet']+'!'+rule['table_name']
        span = target['capabilities'].get('table_ranges', {}).get(identity)
        if not span: fail('table_missing', '模板表对象不存在：' + identity)
        left, top, right, _ = range_boundaries(span)
        if top != tr: fail('table_header_mismatch', '表对象表头与字段定位行不同：' + identity)
        contract['table_ranges'] = {identity: f'{get_column_letter(left)}{top}:{get_column_letter(right)}{max(r["target_row"] for r in locations)}'}
    return {'ok': True, 'operation': 'bind_relationship', 'contract': contract,
        'conditions': {'source_count': 1, 'source_formats': ['xlsx'], 'target_format': 'xlsx',
            'source_guards': [{'source_index': 0, 'ref': sh[k['source_label']]['ref'], 'text': k['source_label']} for k in rule['keys']],
            'target_guards': [{'ref': th[k['target_label']]['ref'], 'text': k['target_label']} for k in rule['keys']]},
        'source_sha256': [source['sha256']], 'target_sha256': target['sha256'],
        'locations': locations, 'source_header_row': sr, 'target_header_row': tr,
        'record_count': len(locations), 'field_count': len(requirements)}
