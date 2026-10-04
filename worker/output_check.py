"""Check a saved artifact against a scoped task contract and current evidence.

The caller supplies coverage separately from the successfully written fields.
This checks explicit evidence/calculations, not unknown business semantics.
"""
import io
import re
from copy import copy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from worker.office_worker import WorkerError, _decode_b64, _sha256
from worker.structured_worker import _read


def _number(value):
    if value is None or isinstance(value, bool):
        raise ValueError('缺失或布尔值不能默认成0')
    try:
        number = Decimal(str(value).replace(',', '').strip())
    except InvalidOperation as exc:
        raise ValueError('来源不是明确数值') from exc
    if not number.is_finite():
        raise ValueError('非有限数值')
    return number


def ratio_tolerance(requirement, fmt):
    raw = requirement.get('absolute_tolerance')
    if 'absolute_tolerance' not in requirement: return Decimal(0)
    if fmt != 'xlsx' or requirement.get('operation') not in {'ratio','percent'} or 'decimals' in requirement or isinstance(raw,bool) or not isinstance(raw,(int,float)):
        raise ValueError('绝对误差仅用于Excel未舍入比例，金额及文字不能使用')
    tolerance = _number(raw)
    if not Decimal(0) < tolerance <= Decimal('1e-12'):
        raise ValueError('比例绝对误差须大于0且不超过1e-12')
    return tolerance


def _expected(requirement, nodes, template_nodes=None):
    op = requirement.get('operation', 'copy')
    if op == 'compose':
        from worker.text_composition import compose_expected
        return compose_expected(requirement, nodes, template_nodes or {}, _expected)
    if op == 'blank':
        from worker.value_rules import blank_expected
        return blank_expected(requirement,nodes)
    if bool(requirement.get('source_unit'))!=bool(requirement.get('target_unit')):
        raise ValueError('来源和目标单位需成对明确；单位依据可来自已确认任务或本次文件，不猜')
    for key,unit_key in [('source_unit_ref','source_unit'),('target_unit_source_ref','target_unit')]:
        if requirement.get(key):
            evidence=nodes.get(requirement[key])
            if not evidence or evidence.get('formula') or evidence.get('value')!=requirement.get(unit_key):
                raise ValueError('单位证据与换算规则不一致')
    refs = requirement.get('source_refs', [])
    if requirement.get('source_matches'):
        if refs:raise ValueError('固定来源引用与按键查找二选一，不混用')
        from worker.source_match import source_match
        refs=[source_match(spec,nodes,template_nodes or {}) for spec in requirement['source_matches']]
    if not isinstance(refs, list) or not refs:
        raise ValueError('必须有本次来源引用；本版不独立证明无来源项目确实缺失')
    if len(set(refs))!=len(refs):raise ValueError('同一来源不能重复计入算式')
    labels=requirement.get('source_labels')
    label_refs=requirement.get('source_label_refs')
    if label_refs is not None and (not isinstance(label_refs,list) or len(label_refs)!=len(refs)):
        raise ValueError('source_label_refs需按来源顺序逐项对应')
    if labels is not None and (not isinstance(labels,list) or len(labels)!=len(refs) or not all(isinstance(v,str) and v for v in labels)):
        raise ValueError('source_labels需按来源顺序逐项列直接字段标签，数量与来源引用一致')
    if requirement.get('source_label') and len(refs)>1:
        raise ValueError('多个来源参与计算，不能要求每个来源都叫同一个派生结果名称；删除source_label，改用source_labels逐项标签，或由各source_matches声明source_label。不改正确产物。')
    values = []
    for index,ref in enumerate(refs):
        node = nodes.get(ref)
        if node is None:
            raise ValueError('来源引用不存在或超出本次读取范围')
        if node.get('formula') or node.get('value_error'):
            raise ValueError('来源公式未计算或类型有错误')
        if node.get('kind')=='filename_identity' and not requirement.get('filename_identity'):
            raise ValueError('文件名仅在任务明确允许身份线索且清单声明filename_identity时使用')
        if node.get('kind')=='filename_identity' and (op not in {'copy','extract'} or requirement.get('source_unit') or requirement.get('decimals') is not None):
            raise ValueError('文件名身份线索不参与金额、单位换算或财务算式')
        if node.get('value') is None:raise ValueError('来源值缺失，不能以目标也为空判完成')
        if node.get('value_type') in {'datetime', 'time'} and (op != 'copy' or len(refs) != 1 or requirement.get('source_unit') or requirement.get('decimals') is not None):
            raise ValueError('日期或时间只支持原值复制，不能参与金额和数值算式')
        label=labels[index] if labels else requirement.get('source_label')
        if not label and requirement.get('source_matches'):label=requirement['source_matches'][index].get('source_label')
        if label and label_refs:
            from worker.value_rules import explicit_label
            explicit_label(node,label,label_refs[index],nodes,ref)
        elif label and label not in [*(node.get('row_headers', [])[:1]),
                                   *(node.get('column_headers', [])[:1])]:
            raise ValueError('来源主标签不符；不能以历史相邻标签强行匹配')
        values.append(node.get('value'))
    if op == 'extract':
        if len(values)!=1:raise ValueError('文本拆分需要恰好一个明确来源')
        from worker.value_rules import extract_text
        value=extract_text(values[0],requirement.get('text_extract'))
        if not requirement.get('source_unit') and requirement.get('decimals') is None:return value
        values=[value];op='copy'
    if op == 'copy' and len(values) == 1 and not requirement.get('source_unit') and requirement.get('decimals') is None:
        return values[0]
    numbers = [_number(v) for v in values]
    if op == 'copy' and len(numbers) == 1:
        value = numbers[0]
    elif op == 'sum':
        value = sum(numbers, Decimal(0))
    elif op == 'product':
        value=Decimal(1)
        for number in numbers: value*=number
    elif op == 'difference' and len(numbers) == 2:
        value = numbers[0] - numbers[1]
    elif op in {'ratio', 'percent'} and len(numbers) == 2:
        if numbers[1] == 0:
            raise ValueError('分母为0，不能默认结果0')
        value = numbers[0] / numbers[1]
        if op == 'percent':
            value *= 100
    elif op == 'formula':
        from worker.value_rules import evaluate_expression
        value=evaluate_expression(requirement.get('expression'),numbers)
    else:
        raise ValueError('未支持或不明确的算式')
    units = {'元': Decimal(1), '千元': Decimal(1000), '万元': Decimal(10000), '亿元': Decimal(100000000)}
    source_unit, target_unit = requirement.get('source_unit'), requirement.get('target_unit')
    if source_unit or target_unit:
        if source_unit not in units or target_unit not in units:
            raise ValueError('单位换算依据不明确')
        value = value * units[source_unit] / units[target_unit]
    places = requirement.get('decimals')
    if places is not None:
        if isinstance(places, bool) or not isinstance(places, int) or not 0 <= places <= 10:
            raise ValueError('小数位需0至10整数')
        value = value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    return value


def _coverage(spec, fmt):
    result = []
    for item in spec:
        if isinstance(item,str) and fmt=='xlsx' and '!' in item and ':' in item.rsplit('!',1)[1]:
            sheet,span=item.rsplit('!',1)
            expanded=_coverage([{'sheet':sheet,'range':span}],fmt)
            result.extend(sorted(expanded))
        elif isinstance(item, str):
            result.append(item)
        elif fmt == 'xlsx' and isinstance(item, dict):
            from openpyxl.utils.cell import range_boundaries, get_column_letter
            sheet, span = item.get('sheet'), item.get('range')
            if not isinstance(sheet, str) or not isinstance(span, str):
                raise WorkerError('invalid_coverage', '覆盖范围需工作表和矩形范围')
            left, top, right, bottom = range_boundaries(span)
            if not all(isinstance(v, int) for v in [left, top, right, bottom]):
                raise WorkerError('invalid_coverage', '覆盖范围不能是整列或整行')
            if (right-left+1)*(bottom-top+1) > 10000:
                raise WorkerError('coverage_limit', '单范围最多10000格')
            result.extend(f'{sheet}!{get_column_letter(c)}{r}'
                          for r in range(top, bottom+1) for c in range(left, right+1))
        else:
            raise WorkerError('invalid_coverage', '覆盖范围格式不支持')
    if not result or len(result) > 10000 or len(set(result)) != len(result):
        raise WorkerError('invalid_coverage', '需非空、无重复且最多10000项覆盖范围')
    return set(result)


def _xlsx_format(template, artifact, coverage, table_ranges):
    from openpyxl import load_workbook
    before = load_workbook(io.BytesIO(template), data_only=False, keep_links=False)
    after = load_workbook(io.BytesIO(artifact), data_only=False, keep_links=False)
    errors = []
    try:
        if before.sheetnames != after.sheetnames:
            return ['工作表名称或顺序变化']
        for name in before.sheetnames:
            old, new = before[name], after[name]
            if str(old.freeze_panes) != str(new.freeze_panes): errors.append(name+': 冻结窗格变化')
            if list(map(str, old.merged_cells.ranges)) != list(map(str, new.merged_cells.ranges)):
                errors.append(name+': 合并区域变化')
            if old.protection != new.protection: errors.append(name+': 工作表保护设置变化')
            for key, dim in old.column_dimensions.items():
                other = new.column_dimensions[key]
                if (dim.width, dim.hidden) != (other.width, other.hidden):
                    errors.append(name+': 列宽或隐藏状态变化 '+key)
            for key, dim in old.row_dimensions.items():
                other = new.row_dimensions[key]
                if (dim.height, dim.hidden) != (other.height, other.hidden):
                    errors.append(name+': 行高或隐藏状态变化 '+str(key))
            for cell in list(old._cells.values()):
                    other = new[cell.coordinate]
                    if any(copy(getattr(cell, p)) != copy(getattr(other, p))
                           for p in ['font', 'fill', 'border', 'alignment', 'protection']):
                        errors.append(name+': 样式变化 '+cell.coordinate)
                    if cell.number_format != other.number_format:
                        errors.append(name+': 数字格式变化 '+cell.coordinate)
                    if name+'!'+cell.coordinate not in coverage and cell.value != other.value:
                        errors.append(name+': 未要求修改的内容变化 '+cell.coordinate)
            for cell in list(new._cells.values()):
                    if name+'!'+cell.coordinate not in coverage and cell.value != old[cell.coordinate].value:
                        message=name+': 未要求修改的内容变化 '+cell.coordinate
                        if message not in errors: errors.append(message)
            if str(old.data_validations)!=str(new.data_validations): errors.append(name+': 数据验证规则变化')
            if str(old.print_area)!=str(new.print_area) or old.print_title_rows!=new.print_title_rows or old.print_title_cols!=new.print_title_cols:
                errors.append(name+': 打印区域或标题变化')
            if old.sheet_properties!=new.sheet_properties or old.page_setup!=new.page_setup or old.page_margins!=new.page_margins:
                errors.append(name+': 页面或工作表设置变化')
            if set(old.tables) != set(new.tables):
                errors.append(name+': 表对象名称变化')
            for table_name in old.tables:
                if table_name not in new.tables: continue
                table = new.tables[table_name]
                expected = table_ranges.get(name+'!'+table_name, old.tables[table_name].ref)
                if table.ref != expected: errors.append(name+': 表范围错误 '+table_name)
                if old.tables[table_name].autoFilter is not None and table.autoFilter is None:
                    errors.append(name+': 表筛选被删除 '+table_name)
                if table.autoFilter and table.autoFilter.ref != expected:
                    errors.append(name+': 表筛选范围错误 '+table_name)
        return errors
    finally:
        before.close(); after.close()


def _xlsx_parts(template, artifact, coverage, table_ranges):
    import zipfile
    from worker.xlsx_package_patch import _sheet_parts, _xml
    before_members=set(_check_zip_names(template))
    after_members=set(_check_zip_names(artifact))
    errors=[]
    if before_members!=after_members:
        errors.append('文件部件集合变化：删除'+str(sorted(before_members-after_members))+
                      '；新增'+str(sorted(after_members-before_members)))
    with zipfile.ZipFile(io.BytesIO(template)) as before, zipfile.ZipFile(io.BytesIO(artifact)) as after:
        sheets=_sheet_parts(before)
        allowed={sheets[ref.rsplit('!',1)[0]] for ref in coverage if ref.rsplit('!',1)[0] in sheets}
        for part in before_members:
            if part.startswith('xl/tables/') and part.endswith('.xml'):
                table=_xml(before.read(part))
                if any(identity.rsplit('!',1)[1]==table.get('name') for identity in table_ranges):
                    allowed.add(part)
        for part in sorted((before_members & after_members)-allowed):
            if before.read(part)!=after.read(part):errors.append('未授权文件部件变化：'+part)
    return errors


def _check_zip_names(data):
    from worker.office_worker import _check_zip
    return _check_zip(data,'output')


def check_output(request):
    items = request.get('sources')
    contract = request.get('contract')
    artifact_item, template_item = request.get('artifact'), request.get('target')
    if not isinstance(items, list) or not 1 <= len(items) <= 8 or not isinstance(contract, dict):
        raise WorkerError('invalid_check_request', '检查需1至8份来源和明确任务清单')
    if not isinstance(contract.get('table_ranges',{}),dict):
        raise WorkerError('invalid_table_ranges','table_ranges需对象，例如{"企业资料!企业资料表":"A1:H5"}，不是数组')
    if not isinstance(artifact_item, dict) or not isinstance(template_item, dict):
        raise WorkerError('invalid_check_request', '检查需当前产物及模板')
    fmt = artifact_item.get('format')
    if fmt not in {'docx', 'xlsx'} or fmt != template_item.get('format'):
        raise WorkerError('invalid_check_format', '产物需与模板同为DOCX或XLSX')
    if contract.get('preserve_only') is True:
        if (contract.get('coverage') != [] or contract.get('requirements') != []
                or contract.get('preserve_template') is not True
                or contract.get('tables') is not None or contract.get('table_ranges')
                or contract.get('allow_existing_values')):
            raise WorkerError('invalid_preserve_only', '整份保留需空填写清单且开启模板保留')
        from worker.evidence_read import read_evidence
        inputs = [read_evidence(item, i, contract) for i, item in enumerate(items)]
        artifact, original = _read(artifact_item, 'source'), _read(template_item, 'source')
        unchanged = artifact['bytes'] == original['bytes']
        read_complete = all(s['capabilities'].get('read') for s in inputs)
        return {'ok': True, 'operation': 'check_output', 'complete': unchanged and read_complete,
                'artifact_sha256': _sha256(artifact['bytes']), 'target_sha256': _sha256(original['bytes']),
                'source_sha256': [s['sha256'] for s in inputs], 'checked': 0, 'passed': 0, 'checks': [],
                'preserve_only': True, 'coverage_errors': [] if read_complete else [{'kind':'unread_source'}],
                'format_errors': [] if unchanged else [{'kind':'preserve_only_bytes_changed'}],
                'scope': '本文件整份原样保留，未新增填写字段；缺资料及人工手续另列于整组任务记录'}
    coverage = _coverage(contract.get('coverage', []), fmt)
    requirements = contract.get('requirements')
    if not isinstance(requirements, list) or not requirements or len(requirements) > 10000:
        raise WorkerError('invalid_requirements', '需1至10000项明确要求')
    from worker.evidence_read import read_evidence
    inputs = [read_evidence(item,i,contract) for i,item in enumerate(items)]
    try:
        artifact = _read(artifact_item, 'source')
    except WorkerError as exc:
        explanation = ('候选输出含重复单元格坐标，请修复这个候选文件里的重复格；原来源文件保持不变'
                       if exc.code == 'duplicate_cell_reference' else '候选输出读取失败：' + exc.message)
        raise WorkerError(exc.code, explanation, {**exc.details, 'file_role': 'candidate_output',
                          'filename': artifact_item.get('filename')}) from exc
    original_target=_read(template_item,'target',target_mode=
        'review_values' if contract.get('allow_existing_values') else 'fill_blanks') if fmt=='docx' else _read(template_item,'source')
    template = _decode_b64(template_item.get('bytes_b64'), 'target')
    data = artifact['bytes']
    sources = {f's{i}::{n["ref"]}': n for i, s in enumerate(inputs) for n in s['nodes']}
    sources.update({f's{i}::$read_scope':s['read_scope_meta'] for i,s in enumerate(inputs)})
    outputs = {n['ref']: n for n in artifact['nodes']}
    template_nodes={n['ref']:n for n in original_target['nodes']}
    if fmt=='docx':
        for slot in original_target['slots']:
            if slot.get('write_kind') == 'text_span':
                start, end = slot['span_start'], slot['span_end']
                template_nodes[slot['target_ref']] = {
                    **slot, 'value': slot['value'][start:end],
                }
        from worker.docx_output_check import docx_output_nodes
        outputs.update(docx_output_nodes(original_target,artifact,coverage))
    if any(r.get('source_matches') for r in requirements) and any(s['capabilities'].get('next_offset') is not None for s in inputs):
        raise WorkerError('incomplete_match_source','按键唯一匹配需要读完相关来源，不能以第一页证明没有重复')
    checks, checked_targets, errors = [], set(), []
    for req in requirements:
        if not isinstance(req, dict) or not isinstance(req.get('target_ref'), str):
            raise WorkerError('invalid_requirement', '每项要求需目标引用')
        ref = req['target_ref']
        if ref in checked_targets: raise WorkerError('duplicate_requirement', '重复目标要求 '+ref)
        checked_targets.add(ref)
        expected=None;actual=None
        try:
            ratio_tolerance(req,fmt)
            expected = _expected(req, sources,template_nodes)
            if req.get('target_unit_ref'):
                unit_node=outputs.get(req['target_unit_ref'])
                if not unit_node or unit_node.get('value')!=req.get('target_unit'):
                    raise ValueError('目标单位引用 '+req['target_unit_ref']+' 实际为 '+repr(unit_node.get('value') if unit_node else None)
                                     +'，要求为 '+repr(req.get('target_unit'))+'；应引用真实单位值而非栏目标题；不需要该证据时修正清单，不修改正确产物')
            if expected is None and fmt=='xlsx':
                from worker.source_xlsx import position
                if '!' not in ref:raise ValueError('空值目标需明确工作表和物理格')
                sheet,coordinate=ref.rsplit('!',1)
                if sheet not in [s['sheet'] for s in artifact['capabilities'].get('sheet_inventory',[])]:
                    raise ValueError('空值目标工作表不存在')
                position(coordinate)
            if ref not in outputs and (expected is not None or fmt=='docx'): raise ValueError('产物未读到该位置')
            actual = outputs.get(ref,{}).get('value')
            if outputs.get(ref,{}).get('formula'): raise ValueError('产物公式未重新计算')
            if outputs.get(ref,{}).get('value_error'): raise ValueError('产物值类型或日期有错误')
            typed_refs = req.get('source_refs', [])
            if req.get('source_matches'):
                from worker.source_match import source_match
                typed_refs = [source_match(spec, sources, template_nodes) for spec in req['source_matches']]
            expected_type = sources.get(typed_refs[0], {}).get('value_type') if req.get('operation', 'copy') == 'copy' and len(typed_refs) == 1 else None
            actual_type = outputs.get(ref, {}).get('value_type')
            if expected is not None and (expected_type or actual_type) and expected_type != actual_type:
                raise ValueError('日期/时间类型与来源不一致，文字外观不能代替Excel日期')
            numeric = isinstance(expected, Decimal)
            if numeric and fmt=='xlsx' and (isinstance(actual,bool) or not isinstance(actual,(int,float))):
                raise ValueError('Excel计算金额必须是数值，不能写成数字外观的文本')
            if numeric:
                number_text=actual[:-1] if fmt=='docx' and req.get('render')=='percent_text' and isinstance(actual,str) and actual.endswith('%') else actual
                tolerance = ratio_tolerance(req,fmt)
                correct = abs(_number(number_text)-expected) <= tolerance
                if fmt=='docx' and req.get('render'):
                    places=req.get('decimals')
                    if not isinstance(places,int) or isinstance(places,bool) or not 0<=places<=10:
                        raise ValueError('Word固定小数格式需要0至10整数decimals')
                    if req['render'] not in {'fixed_decimal','percent_text'}:
                        raise ValueError('未支持的Word数值显示方式')
                    suffix='%' if req['render']=='percent_text' else ''
                    correct = actual==format(expected,f'.{places}f')+suffix
            else:
                if expected is None:
                    correct = actual in (None,'')
                else:
                    correct = (actual==str(expected)) if fmt=='docx' else type(actual) is type(expected) and actual == expected
            if not correct: raise ValueError('产物值与当前来源或明确算式不一致')
            subject = req.get('subject_source_ref')
            if subject:
                subject_node = sources.get(subject)
                target_key = outputs.get(req.get('subject_target_ref'))
                if not subject_node or not target_key or subject_node['value'] != target_key['value']:
                    raise ValueError('该值所属主体与目标记录不一致')
            checks.append({'target_ref': ref, 'passed': True})
        except (ValueError, InvalidOperation) as exc:
            checks.append({'target_ref': ref, 'passed': False, 'reason': str(exc),
                           'expected':str(expected) if isinstance(expected,Decimal) else expected,
                           'actual':outputs.get(ref,{}).get('value'), 'source_refs':req.get('source_refs',[])})
    if coverage != checked_targets:
        errors.append({'kind': 'coverage', 'missing': sorted(coverage-checked_targets),
                       'outside': sorted(checked_targets-coverage),
                       'instruction':'coverage只列任务需核对的数据位置，必须与完整requirements一致；保留模板表头无需加入coverage。不能只按已经写对的项缩小业务要求。'})
    format_errors = []
    if contract.get('preserve_template', True):
        if fmt == 'xlsx':
            format_errors = _xlsx_format(template, data, coverage, contract.get('table_ranges', {}))
            format_errors += _xlsx_parts(template,data,coverage,contract.get('table_ranges',{}))
        else:
            from worker.docx_output_check import check_docx_format
            format_errors = check_docx_format(template,data,coverage)
    complete_read = all(s['capabilities'].get('read') for s in inputs)
    if not complete_read: errors.append({'kind': 'unread_source'})
    passed = sum(c['passed'] for c in checks)
    complete = passed == len(checks) and not errors and not format_errors
    return {'ok': True, 'operation': 'check_output', 'artifact_sha256': _sha256(data),
            'source_sha256': [s['sha256'] for s in inputs], 'target_sha256': _sha256(template),
            'complete': complete, 'checked': len(checks), 'passed': passed, 'checks': checks,
            'coverage_errors': errors, 'format_errors': format_errors,
            'scope': '仅核对独立列明范围及明确引用/算式；不保证清单理解了所有陌生业务，也未验证打印分页'}
