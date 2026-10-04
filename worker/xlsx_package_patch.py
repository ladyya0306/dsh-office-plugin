"""Replace authorized worksheet values without reserializing other ZIP parts."""
from __future__ import annotations

from collections import defaultdict
import io
import math
import posixpath
import re
import zipfile

from worker.package_writer import replace_parts
_ETREE = None

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
CELL = re.compile(r"^([A-Z]{1,3})([1-9][0-9]{0,6})$")


class XlsxPackageError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code, self.message = code, message


def _etree():
    global _ETREE
    if _ETREE is None:
        try:
            from lxml import etree
        except ModuleNotFoundError as exc:
            raise XlsxPackageError("dependency_unavailable", "XLSX最小部件写入需要受信任解释器中的lxml") from exc
        _ETREE = etree
    return _ETREE


def _xml(data: bytes):
    etree = _etree()
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    root = etree.fromstring(data, parser)
    if root.getroottree().docinfo.doctype:
        raise XlsxPackageError("unsupported_xml_doctype", "工作表部件包含DTD，拒绝修改")
    return root


def _column(coordinate: str) -> int:
    match = CELL.fullmatch(coordinate)
    if not match:
        raise XlsxPackageError("invalid_cell_reference", "单元格引用不合法")
    number = 0
    for letter in match[1]:
        number = number * 26 + ord(letter) - ord("A") + 1
    return number


def _sheet_parts(archive) -> dict[str, str]:
    workbook = _xml(archive.read("xl/workbook.xml"))
    relationships = _xml(archive.read("xl/_rels/workbook.xml.rels"))
    rels = {}
    for item in relationships.findall(f"{{{PKG_REL}}}Relationship"):
        identity = item.get("Id")
        if identity in rels:
            raise XlsxPackageError("duplicate_relationship", "工作簿关系ID重复")
        rels[identity] = item
    parts = {}
    for sheet in workbook.findall(f"{{{MAIN}}}sheets/{{{MAIN}}}sheet"):
        relationship = rels.get(sheet.get(f"{{{REL}}}id"))
        if relationship is None or relationship.get("TargetMode") == "External" or not relationship.get("Type", "").endswith("/worksheet"):
            raise XlsxPackageError("unsupported_sheet_relationship", "工作表关系缺失或类型不受支持")
        target = relationship.get("Target", "")
        part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
        if not part.startswith("xl/") or part not in archive.namelist():
            raise XlsxPackageError("invalid_sheet_relationship", "工作表关系越界或部件缺失")
        name = sheet.get("name")
        if name in parts:
            raise XlsxPackageError("duplicate_sheet_name", "工作表名称重复")
        parts[name] = part
    return parts


def _set_cell(sheet, coordinate: str, value):
    etree = _etree()
    if sheet.tag != f"{{{MAIN}}}worksheet":
        raise XlsxPackageError("unsupported_sheet_xml", "工作表XML类型不受支持")
    row_number = int(CELL.fullmatch(coordinate)[2])
    data = sheet.find(f"{{{MAIN}}}sheetData")
    if data is None:
        raise XlsxPackageError("missing_sheet_data", "工作表没有单元格数据区域")
    rows = [row for row in data if row.tag == f"{{{MAIN}}}row" and row.get("r") == str(row_number)]
    if len(rows) > 1:
        raise XlsxPackageError("duplicate_row_reference", "工作表行引用重复")
    row = rows[0] if rows else etree.Element(f"{{{MAIN}}}row", r=str(row_number))
    if not rows:
        at = next((i for i, child in enumerate(data) if child.tag == row.tag and int(child.get("r")) > row_number), len(data))
        data.insert(at, row)
    cells = [cell for cell in row if cell.tag == f"{{{MAIN}}}c" and cell.get("r") == coordinate]
    if len(cells) > 1:
        raise XlsxPackageError("duplicate_cell_reference", "单元格引用重复")
    cell = cells[0] if cells else etree.Element(f"{{{MAIN}}}c", r=coordinate)
    if not cells:
        column = _column(coordinate)
        at = next((i for i, child in enumerate(row) if child.tag == cell.tag and _column(child.get("r", "")) > column), len(row))
        row.insert(at, cell)
    if cell.find(f"{{{MAIN}}}f") is not None:
        raise XlsxPackageError("formula_target", "禁止改写公式单元格")
    for child in list(cell):
        if child.tag in {f"{{{MAIN}}}v", f"{{{MAIN}}}is"}:
            cell.remove(child)
    if value is None:
        cell.attrib.pop('t',None)
    elif isinstance(value, str):
        cell.set("t", "inlineStr")
        inline = etree.Element(f"{{{MAIN}}}is")
        text = etree.SubElement(inline, f"{{{MAIN}}}t")
        text.set(XML_SPACE, "preserve")
        text.text = value
        cell.insert(0, inline)
    elif isinstance(value, bool):
        cell.set("t", "b")
        item = etree.Element(f"{{{MAIN}}}v"); item.text = "1" if value else "0"; cell.insert(0, item)
    elif isinstance(value, (int, float)) and math.isfinite(value):
        cell.set("t", "n")
        item = etree.Element(f"{{{MAIN}}}v"); item.text = str(value); cell.insert(0, item)
    else:
        raise XlsxPackageError("unsupported_value_type", "仅支持已校验的文本、数值、布尔值或明确空值")


def patch_xlsx(target: bytes, changes: list[dict]) -> bytes:
    """Input must already pass worker feature, size, reference and value checks."""
    with zipfile.ZipFile(io.BytesIO(target)) as archive:
        infos = archive.infolist()
        if len({info.filename for info in infos}) != len(infos):
            raise XlsxPackageError("duplicate_package_member", "ZIP部件名称重复")
        parts = _sheet_parts(archive)
        grouped = defaultdict(list)
        seen = set()
        for change in changes:
            sheet, coordinate = change["target_ref"].rsplit("!", 1)
            if sheet not in parts or not CELL.fullmatch(coordinate):
                raise XlsxPackageError("invalid_target_reference", "目标工作表或单元格不存在")
            identity = (sheet, coordinate)
            if identity in seen:
                raise XlsxPackageError("duplicate_target", "禁止重复目标位置")
            seen.add(identity); grouped[parts[sheet]].append((coordinate, change["value"]))
        replacements = {}
        for part, edits in grouped.items():
            sheet = _xml(archive.read(part))
            for coordinate, value in edits:
                _set_cell(sheet, coordinate, value)
            dimension=sheet.find(f'{{{MAIN}}}dimension')
            if dimension is not None:
                from openpyxl.utils.cell import range_boundaries,get_column_letter
                left,top,right,bottom=range_boundaries(dimension.get('ref',''))
                for coordinate,_value in edits:
                    row=int(CELL.fullmatch(coordinate)[2]);column=_column(coordinate)
                    left=min(left,column);top=min(top,row);right=max(right,column);bottom=max(bottom,row)
                dimension.set('ref',f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}')
            replacements[part] = _etree().tostring(sheet, xml_declaration=True, encoding="UTF-8", standalone=True)
        return replace_parts(target, replacements)
