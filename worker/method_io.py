"""Trusted Office callbacks. Never expose Python host objects to method code."""
import base64
import copy
import io
import json
import re
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from worker.office_worker import WorkerError, _decode_b64, _sha256, _check_zip
from worker.structured_worker import _read
from worker.xlsx_package_patch import patch_xlsx, _xml, XlsxPackageError
from worker.package_writer import replace_parts
from worker.output_check import _coverage


def plain(value):
    # The sandbox boundary gets only JSON values, never functions/Office objects.
    text = json.dumps(value, ensure_ascii=False, allow_nan=False)
    if len(text.encode('utf-8')) > 2*1024*1024:
        raise WorkerError('method_data_limit', '方法数据超过2MiB')
    return json.loads(text)


def output_name(name, fmt):
    if (not isinstance(name, str) or not 1 <= len(name) <= 120 or name != name.strip()
        or re.search(r'[\\/:*?"<>|\x00-\x1f]', name) or not name.lower().endswith('.'+fmt)
        or re.fullmatch(r'(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])', name.split('.',1)[0], re.I)):
        raise WorkerError('method_output_name', '输出必须是当前工作区普通新文件名')
    return name


def _patch_tables(data, table_ranges):
    if not table_ranges: return data
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=False, keep_links=False)
    known = {sheet.title+'!'+t.name: (t.name, t.id)
             for sheet in book for t in sheet.tables.values()}
    book.close()
    if set(table_ranges)-set(known):
        raise WorkerError('unknown_table', '表对象不存在')
    replacements = {}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        table_parts = {name: _xml(archive.read(name)) for name in archive.namelist()
                       if name.startswith('xl/tables/') and name.endswith('.xml')}
        for identity, span in table_ranges.items():
            if not isinstance(span,str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*:[A-Z]{1,3}[1-9][0-9]*',span):
                raise WorkerError('invalid_table_range','表范围需明确矩形')
            matches = [(part, tree) for part, tree in table_parts.items()
                       if tree.get('name')==known[identity][0] and tree.get('id')==str(known[identity][1])]
            if len(matches)!=1: raise WorkerError('ambiguous_table','表部件身份不唯一')
            part, tree=matches[0]; tree.set('ref',span)
            for child in tree:
                if child.tag.endswith('}autoFilter'): child.set('ref',span)
            from lxml import etree
            replacements[part]=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    return replace_parts(data,replacements)


class OfficeCallbacks:
    def __init__(self, request, evidence_cache=None):
        from worker.evidence_read import read_evidence
        self.source_items=request['sources']
        self.sources = [read_evidence(item,i,request['contract'],request.get('conditions',{}).get('source_guards',[]),
                                      cache=evidence_cache)
                        for i,item in enumerate(request['sources'])]
        self.target = _read(request['target'],'target',target_mode=
                            'review_values' if request['contract'].get('allow_existing_values') else 'fill_blanks')
        self.outputs = {}
        self.contract=request['contract']
        if self.contract.get('preserve_only') is True:
            if (self.contract.get('coverage') != [] or self.contract.get('requirements') != []
                    or self.contract.get('preserve_template') is not True
                    or self.contract.get('tables') is not None or self.contract.get('table_ranges')
                    or self.contract.get('allow_existing_values')):
                raise WorkerError('invalid_preserve_only','整份保留不得同时声明填写范围')
            self.allowed=set()
        else:
            self.allowed=_coverage(self.contract.get('coverage',[]),self.target['format'])
        self.required={r['target_ref'] for r in self.contract.get('requirements',[])}

    def check_conditions(self, conditions):
        if not isinstance(conditions,dict) or conditions.get('source_count')!=len(self.sources):
            raise WorkerError('method_not_applicable','来源数量不同；回到DSH识别本次任务，不强用旧方法')
        formats=conditions.get('source_formats')
        if formats!=[s['format'] for s in self.sources] or conditions.get('target_format')!=self.target['format']:
            raise WorkerError('method_not_applicable','来源或目标格式不同')
        for guard in conditions.get('source_guards',[]):
            index=guard.get('source_index')
            if isinstance(index,bool) or not isinstance(index,int) or not 0<=index<len(self.sources):
                raise WorkerError('invalid_method_guard','来源结构条件缺少有效编号')
            self._guard(self.sources[index],guard)
        for guard in conditions.get('target_guards',[]): self._guard(self.target,guard)

    @staticmethod
    def _guard(document,guard):
        found=next((n for n in document['nodes'] if n['ref']==guard.get('ref')),None)
        if not found or found.get('formula') or found.get('value')!=guard.get('text'):
            raise WorkerError('method_not_applicable','结构或标签不同：'+str(guard.get('ref'))+'；交给DSH重新判断')

    def read_source(self, index, requested_fields=None, source_offset=0, read_scope='evidence'):
        if isinstance(index,bool) or not isinstance(index,int) or not 0<=index<len(self.sources):
            raise WorkerError('invalid_source_index','来源编号不属于本次输入')
        if read_scope not in {'evidence', 'full'}:
            raise WorkerError('invalid_read_scope', 'read_scope只能为evidence或full')
        if read_scope == 'full':
            if requested_fields is not None or source_offset != 0:
                raise WorkerError('invalid_read_scope', '整份读取不能同时指定检索词或分页起点')
            from worker.evidence_read import read_evidence
            source = read_evidence(self.source_items[index], index, {'requirements': []})
            return plain({'format': source['format'], 'nodes': source['nodes'], 'capabilities': source['capabilities']})
        # Programs can search labels outside the checklist's selected evidence.
        source=self.sources[index] if requested_fields is None and source_offset==0 else _read(
            self.source_items[index],'source',fields=requested_fields,offset=source_offset)
        return plain({'format':source['format'],'nodes':source['nodes'],
                      'capabilities':source['capabilities']})

    def read_target(self):
        return plain({'format':self.target['format'],'nodes':self.target['nodes'],
                      'target_slots':self.target['slots'],
                      'capabilities':self.target['capabilities']})

    def convert_amount(self, value, source_unit, target_unit, decimals=2):
        units={'元':Decimal(1),'千元':Decimal(1000),'万元':Decimal(10000),'亿元':Decimal(100000000)}
        if source_unit not in units or target_unit not in units or isinstance(value,bool) or value is None:
            raise WorkerError('invalid_conversion','金额或单位不明确；缺失不能为0')
        if isinstance(decimals,bool) or not isinstance(decimals,int) or not 0<=decimals<=10:
            raise WorkerError('invalid_decimals','小数位需0至10整数')
        number=Decimal(str(value))
        if len(str(value))>128 or abs(number.adjusted())>100:
            raise WorkerError('invalid_amount','金额字符串或指数超限')
        if not number.is_finite(): raise WorkerError('invalid_amount','金额必须有限')
        return float((number*units[source_unit]/units[target_unit]).quantize(
            Decimal(1).scaleb(-decimals),rounding=ROUND_HALF_UP))

    def calculate(self, operation, values, decimals=2, source_unit=None, target_unit=None,expression=None):
        if not isinstance(values,list) or not 1<=len(values)<=1000:
            raise WorkerError('invalid_calculation','计算需1至1000个明确数值')
        from worker.output_check import _expected
        from worker.output_check import _number
        for value in values:
            number=_number(value)
            if len(str(value))>128 or abs(number.adjusted())>100:raise WorkerError('invalid_calculation','数值字符串或指数超限')
        nodes={str(i):{'value':v} for i,v in enumerate(values)}
        rule={'source_refs':list(nodes),'operation':operation,'decimals':decimals}
        if expression is not None:rule['expression']=expression
        if source_unit or target_unit:rule.update(source_unit=source_unit,target_unit=target_unit)
        try:return float(_expected(rule,nodes))
        except (ValueError,ArithmeticError) as exc:raise WorkerError('invalid_calculation',str(exc)) from exc

    def calculate_many(self, calculations):
        if not isinstance(calculations,list) or not 1<=len(calculations)<=1000:
            raise WorkerError('invalid_calculation','批量计算需1至1000项')
        results=[]
        for calculation in calculations:
            if not isinstance(calculation,dict) or set(calculation)-{'operation','values','decimals','source_unit','target_unit','expression'}:
                raise WorkerError('invalid_calculation','批量计算项只接受operation、values、小数位和单位')
            results.append(self.calculate(**calculation))
        return plain(results)

    def write_cells(self, changes, name, table_ranges=None):
        fmt=self.target['format']; output_name(name,fmt)
        if name in self.outputs: raise WorkerError('duplicate_output','本次输出名重复')
        if len(self.outputs)>=16: raise WorkerError('method_output_limit','一次方法最多生成16份文件')
        if not isinstance(changes,list) or not 1<=len(changes)<=10000:
            raise WorkerError('method_change_limit','本次需1至10000项变更')
        for item in changes:
            if not isinstance(item,dict) or set(item)!={'target_ref','value'}:
                raise WorkerError('invalid_method_change','变更只含目标引用及值')
            if not isinstance(item['target_ref'],str): raise WorkerError('invalid_method_change','目标引用需文字')
            if item['target_ref'] not in self.allowed or item['target_ref'] not in self.required:
                raise WorkerError('method_write_outside_contract','写入位置不在独立任务清单内')
        if (table_ranges or {})!=self.contract.get('table_ranges',{}):
            raise WorkerError('method_table_outside_contract','表范围必须来自独立清单，不由代码扩大')
        if fmt=='docx':
            if table_ranges: raise WorkerError('method_table_outside_contract','Word不接受Excel表范围')
            slots={slot['target_ref']:slot for slot in self.target['slots']}
            refs=[item['target_ref'] for item in changes]
            if len(set(refs))!=len(refs):raise WorkerError('duplicate_target','同一Word位置不能重复写入')
            if any(ref not in slots for ref in refs):
                raise WorkerError('method_docx_target_unsupported','位置不属于本次可写Word对象；请由DSH处理支持范围外的结构')
            edited=[]
            for item in changes:
                value=item['value']
                if value is None or isinstance(value,(dict,list,bool)):
                    raise WorkerError('invalid_method_change','Word值需明确文字或数值；缺失需先按清单处理')
                edited.append({'target_ref':item['target_ref'],'value':str(value)})
            from worker.structured_fill import patch_and_verify
            data=patch_and_verify(self.target,edited)
            if len(data)>15*1024*1024:raise WorkerError('method_output_limit','输出超过15MiB')
            _check_zip(data,'output');self.outputs[name]=data
            return {'output_name':name,'sha256':_sha256(data),'changed':len(changes)}
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(self.target['bytes']),keep_links=False)
        try:
            for item in changes:
                sheet,coordinate=item['target_ref'].rsplit('!',1)
                if sheet not in book: raise WorkerError('invalid_target_reference','目标工作表不存在')
                ws=book[sheet];cell=ws[coordinate]
                if ws.protection.sheet and cell.protection.locked:
                    raise WorkerError('protected_target','目标受保护，不能修改')
                if any(coordinate in region and coordinate!=region.start_cell.coordinate for region in ws.merged_cells.ranges):
                    raise WorkerError('merged_target','只能写合并区域的主单元格')
                if cell.value is not None and cell.value!=item['value'] and not self.contract.get('allow_existing_values',False):
                    raise WorkerError('existing_target','清单未明确允许改写已有值')
        finally: book.close()
        try:
            data=patch_xlsx(self.target['bytes'],plain(changes))
            data=_patch_tables(data,plain(table_ranges or {}))
        except XlsxPackageError as exc:
            raise WorkerError(exc.code,exc.message) from exc
        if len(data)>15*1024*1024: raise WorkerError('method_output_limit','输出超过15MiB')
        _check_zip(data,'output')
        self.outputs[name]=data
        return {'output_name':name,'sha256':_sha256(data),'changed':len(changes)}

    def functions(self):
        return {'read_source':self.read_source,'read_target':self.read_target,
                'write_cells':self.write_cells,'convert_amount':self.convert_amount,'calculate':self.calculate,
                'calculate_many':self.calculate_many}
